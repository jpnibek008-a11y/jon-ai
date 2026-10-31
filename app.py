
import streamlit.components.v1 as components
import streamlit as st
import pandas as pd
import requests
import json
from io import BytesIO
from google import genai
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

st.set_page_config(
    page_title="JON AI 3.5",
    page_icon="💎",
    layout="wide"
)

APP_VERSION = "3.5"
MODEL = "gemini-3.5-flash"

JON_CORE = """
Sen JON AI 3.5 nomli aqlli virtual yordamchisan.
O'zbek tilida ravon gapirasan.
Savollarga aniq va tushunarli javob berasan.
Ta'lim, tarjima, dasturlash va ijodiy ishlarda yordam berasan.
Foydalanuvchi faqat personajning tashqi ko'rinishini o'zgartiradi.
Sening asosiy xaraktering va shaxsiyating o'zgarmaydi.
"""

defaults = {
    "messages": [],
    "tasks": [],
    "character_name": "JON",
    "character_style": "Hologramma",
    "avatar": None,
    "api_key": ""
}

for key, value in defaults.items():
    if key not in st.session_state:
        st.session_state[key] = value


def ask_ai(prompt, extra_context=""):
    api_key = st.session_state.get("api_key", "")

    if not api_key:
        return "Sozlamalardan Gemini API kalitini kiriting."

    try:
        client = genai.Client(api_key=api_key)

        response = client.models.generate_content(
            model=MODEL,
            contents=JON_CORE + "\n" + extra_context + "\n" + prompt
        )

        return response.text or "Javob olinmadi."

    except Exception as e:
        return f"Xatolik: {e}"

def show_character():
    import streamlit.components.v1 as components

    html = r"""
    <!DOCTYPE html>
    <html>
    <head>
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <style>
      * { margin: 0; padding: 0; }
      body {
        background: radial-gradient(ellipse at center, #10285c 0%, #050816 75%);
        overflow: hidden;
      }
      #scene {
        width: 100%;
        height: 510px;
        display: block;
      }
      .caption {
        position: absolute;
        bottom: 16px;
        width: 100%;
        text-align: center;
        color: #70eaff;
        font: 14px Arial, sans-serif;
        letter-spacing: 5px;
        text-shadow: 0 0 12px #00d9ff;
        pointer-events: none;
      }
    </style>
    </head>
    <body>
    <div id="scene"></div>
    <div class="caption">J O N   A I</div>

    <script type="module">
    import * as THREE from
    'https://cdn.jsdelivr.net/npm/three@0.160.0/build/three.module.js';

    const container = document.getElementById("scene");

    const scene = new THREE.Scene();

    const camera = new THREE.PerspectiveCamera(
      45, container.clientWidth / container.clientHeight, 0.1, 100
    );
    camera.position.set(0, 0, 8);

    const renderer = new THREE.WebGLRenderer({
      antialias: true,
      alpha: true
    });

    renderer.setPixelRatio(Math.min(window.devicePixelRatio, 2));
    renderer.setSize(container.clientWidth, container.clientHeight);
    renderer.setClearColor(0x000000, 0);
    container.appendChild(renderer.domElement);

    scene.add(new THREE.AmbientLight(0x8abfff, 2));

    const blueLight = new THREE.PointLight(0x00d9ff, 35, 20);
    blueLight.position.set(2, 3, 4);
    scene.add(blueLight);

    const purpleLight = new THREE.PointLight(0x7b2cff, 25, 20);
    purpleLight.position.set(-3, -2, 3);
    scene.add(purpleLight);

    const character = new THREE.Group();
    scene.add(character);

    // Kristall tanasi
    const coreGeometry = new THREE.IcosahedronGeometry(1.15, 2);
    const coreMaterial = new THREE.MeshPhysicalMaterial({
      color: 0x087dff,
      metalness: 0.55,
      roughness: 0.12,
      transmission: 0.25,
      transparent: true,
      opacity: 0.88,
      emissive: 0x003bff,
      emissiveIntensity: 1.2
    });

    const core = new THREE.Mesh(coreGeometry, coreMaterial);
    character.add(core);

    // Ichki energiya yadrosi
    const inner = new THREE.Mesh(
      new THREE.OctahedronGeometry(0.48, 1),
      new THREE.MeshBasicMaterial({
        color: 0x6ffaff,
        wireframe: true
      })
    );
    character.add(inner);

    // Yorqin ko'zlar
    const eyeMaterial = new THREE.MeshBasicMaterial({
      color: 0x9fffff
    });

    function createEye(x, y) {
      const eye = new THREE.Mesh(
        new THREE.SphereGeometry(0.105, 20, 20),
        eyeMaterial
      );
      eye.position.set(x, y, 0.92);
      character.add(eye);
    }

    createEye(-0.34, 0.18);
    createEye(0.34, 0.18);

    // Gologramma halqalari
    const rings = [];

    for (let i = 0; i < 3; i++) {
      const ring = new THREE.Mesh(
        new THREE.TorusGeometry(1.55 + i * 0.18, 0.012, 8, 120),
        new THREE.MeshBasicMaterial({
          color: i === 1 ? 0x9b55ff : 0x00eaff,
          transparent: true,
          opacity: 0.8
        })
      );

      ring.rotation.x = Math.PI / 2 + i * 0.45;
      ring.rotation.y = i * 0.6;
      character.add(ring);
      rings.push(ring);
    }

    // Suzib yuruvchi kristallar
    const crystals = [];

    for (let i = 0; i < 8; i++) {
      const crystal = new THREE.Mesh(
        new THREE.OctahedronGeometry(0.12 + Math.random() * 0.1),
        new THREE.MeshBasicMaterial({
          color: i % 2 ? 0x00eaff : 0x9b55ff,
          wireframe: true
        })
      );

      const angle = i / 8 * Math.PI * 2;
      crystal.position.set(
        Math.cos(angle) * 2.1,
        Math.sin(angle) * 1.5,
        Math.sin(angle * 2) * 0.4
      );

      character.add(crystal);
      crystals.push(crystal);
    }

    // Yulduzsimon zarrachalar
    const particleGeometry = new THREE.BufferGeometry();
    const particlePositions = [];

    for (let i = 0; i < 450; i++) {
      particlePositions.push(
        (Math.random() - 0.5) * 12,
        (Math.random() - 0.5) * 8,
        (Math.random() - 0.5) * 7
      );
    }

    particleGeometry.setAttribute(
      "position",
      new THREE.Float32BufferAttribute(particlePositions, 3)
    );

    const particles = new THREE.Points(
      particleGeometry,
      new THREE.PointsMaterial({
        color: 0x62eaff,
        size: 0.025,
        transparent: true,
        opacity: 0.8
      })
    );

    scene.add(particles);

    // Animatsiya
    const clock = new THREE.Clock();

    function animate() {
      requestAnimationFrame(animate);

      const t = clock.getElapsedTime();

      character.position.y = Math.sin(t * 1.4) * 0.16;
      character.rotation.y = Math.sin(t * 0.5) * 0.35;
      character.rotation.x = Math.sin(t * 0.7) * 0.08;

      core.rotation.y += 0.003;
      core.rotation.x += 0.002;
      inner.rotation.y -= 0.012;
      inner.rotation.x += 0.008;

      rings.forEach((ring, i) => {
        ring.rotation.z += 0.003 + i * 0.001;
      });

      crystals.forEach((crystal, i) => {
        crystal.rotation.x += 0.012;
        crystal.rotation.y += 0.009;
        crystal.position.y += Math.sin(t * 1.5 + i) * 0.0015;
      });

      particles.rotation.y = t * 0.025;

      renderer.render(scene, camera);
    }

    animate();

    window.addEventListener("resize", () => {
      const width = container.clientWidth;
      const height = container.clientHeight;

      camera.aspect = width / height;
      camera.updateProjectionMatrix();
      renderer.setSize(width, height);
    });
    </script>
    </body>
    </html>
    """

    components.html(html, height=520, scrolling=False)
    

        
                

        

            
        </svg>

        <div class="holo-scan"></div>
        <div class="holo-title">JON AI</div>
    </div>
    </body>
    </html>
    """

    components.html(
        html_code,
        height=490,
        scrolling=False
    )    
def parse_slides(result):
    try:
        clean = result.strip()
        clean = clean.replace("```json", "").replace("```", "")
        start = clean.find("[")
        end = clean.rfind("]") + 1

        if start == -1 or end == 0:
            raise ValueError("JSON topilmadi")

        return json.loads(clean[start:end])

    except Exception:
        return [
            {
                "title": "Taqdimot",
                "points": [result]
            }
        ]


def create_presentation(topic, slides_data):
    prs = Presentation()

    for item in slides_data:
        slide = prs.slides.add_slide(prs.slide_layouts[1])

        slide.shapes.title.text = str(
            item.get("title", topic)
        )

        body = slide.placeholders[1].text_frame
        points = item.get("points", [])

        if isinstance(points, str):
            points = [points]

        body.text = ""

        for point in points:
            paragraph = body.add_paragraph()
            paragraph.text = str(point)
            paragraph.level = 0

    output = BytesIO()
    prs.save(output)
    output.seek(0)

    return output


pages = [
    "🏠 Bosh sahifa",
    "🤖 AI Chat",
    "🎓 Ta'lim markazi",
    "📊 PowerPoint",
    "🌐 Tarjimon",
    "🎨 Rasm prompti",
    "📈 Excel tahlil",
    "📝 Vazifalar",
    "📷 Kamera",
    "👤 Personaj",
    "⚙️ Sozlamalar",
    "💎 Premium"
]

st.sidebar.title("💎 JON AI 3.5")
st.sidebar.caption("Virtual Intelligence System")

page = st.sidebar.radio(
    "Bo'limni tanlang",
    pages
)

st.sidebar.divider()
st.sidebar.caption("JON AI 3.5 | 2026")


if page == "🏠 Bosh sahifa":
    st.title("💎 JON AI 3.5")
    st.caption("Sizning shaxsiy virtual yordamchingiz")

    show_character()

    st.markdown("### Xush kelibsiz!")
    st.write(
        "JON AI bilan suhbatlashing, bilim oling, "
        "tarjima qiling va yangi loyihalar yarating."
    )

    if st.button("🤖 AI Chatni ochish"):
        st.session_state.open_chat = True
        st.rerun()


elif page == "🤖 AI Chat":
    st.title("🤖 JON AI — Aqlli suhbat")

    for message in st.session_state.messages:
        with st.chat_message(message["role"]):
            st.markdown(message["content"])

    prompt = st.chat_input("JON AI ga xabar yozing...")

    if prompt:
        st.session_state.messages.append(
            {"role": "user", "content": prompt}
        )

        with st.chat_message("user"):
            st.markdown(prompt)

        with st.chat_message("assistant"):
            with st.spinner("JON AI javob tayyorlamoqda..."):
                answer = ask_ai(prompt)
                st.markdown(answer)

        st.session_state.messages.append(
            {"role": "assistant", "content": answer}
        )


elif page == "🎓 Ta'lim markazi":
    st.title("🎓 JON AI — Ta'lim markazi")

    subject = st.selectbox(
        "Fan yoki yo'nalish",
        [
            "Matematika",
            "Ingliz tili",
            "Fizika",
            "Kimyo",
            "Tarix",
            "Informatika",
            "Psixologiya",
            "Boshqa"
        ]
    )

    question = st.text_area(
        "Savolingiz yoki mavzuni yozing"
    )

    if st.button("📚 Tushuntirish"):
        if question:
            with st.spinner("Tayyorlanmoqda..."):
                answer = ask_ai(
                    question,
                    "Fan: " + subject +
                    ". Mavzuni o'quvchiga sodda tushuntir."
                )
                st.markdown(answer)
        else:
            st.warning("Avval savol yozing.")


elif page == "📊 PowerPoint":
    st.title("📊 JON AI — PowerPoint yaratuvchi")

    topic = st.text_input(
        "Taqdimot mavzusi"
    )

    slide_count = st.slider(
        "Slaydlar soni",
        3,
        15,
        6
    )

    if st.button("✨ Taqdimot yaratish"):
        if topic:
            prompt = (
                "Mavzu: " + topic +
                ". " + str(slide_count) +
                " ta slayd uchun ma'lumot tayyorla. "
                "Faqat JSON massiv qaytar. "
                'Har bir elementda "title" va "points" bo\'lsin. '
                "points ro'yxat ko'rinishida bo'lsin."
            )

            with st.spinner("Taqdimot tayyorlanmoqda..."):
                result = ask_ai(prompt)
                slides_data = parse_slides(result)
                pptx_file = create_presentation(
                    topic,
                    slides_data
                )

                st.success("Taqdimot tayyor!")

                st.download_button(
                    "⬇️ PowerPoint yuklab olish",
                    data=pptx_file,
                    file_name="JON_AI_Taqdimot.pptx",
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
                )
        else:
            st.warning("Mavzuni kiriting.")


elif page == "🌐 Tarjimon":
    st.title("🌐 JON AI — Universal tarjimon")

    source_lang = st.selectbox(
        "Matn qaysi tilda?",
        [
            "Avtomatik aniqlash",
            "O'zbek",
            "English",
            "Русский",
            "Turk",
            "Arab",
            "Nemis",
            "Fransuz"
        ]
    )

    target_lang = st.selectbox(
        "Qaysi tilga tarjima qilinsin?",
        [
            "O'zbek",
            "English",
            "Русский",
            "Turk",
            "Arab",
            "Nemis",
            "Fransuz"
        ]
    )

    text_to_translate = st.text_area(
        "Tarjima qilinadigan matn"
    )

    if st.button("🌐 Tarjima qilish"):
        if text_to_translate:
            prompt = (
                "Quyidagi matnni " + source_lang +
                " tilidan " + target_lang +
                " tiliga tarjima qil. "
                "Faqat tarjimani chiqar:\n" +
                text_to_translate
            )

            with st.spinner("Tarjima qilinmoqda..."):
                translated = ask_ai(prompt)
                st.markdown("### Natija")
                st.write(translated)
        else:
            st.warning("Matn kiriting.")


elif page == "🎨 Rasm prompti":
    st.title("🎨 JON AI — Rasm prompti")

    idea = st.text_area(
        "Qanday rasm yaratmoqchisiz?"
    )

    style = st.selectbox(
        "Rasm uslubi",
        [
            "Realistik",
            "3D",
            "Multfilm",
            "Anime",
            "Kiberpank",
            "Fantastik"
        ]
    )

    if st.button("✨ Prompt yaratish"):
        if idea:
            prompt = (
                "Quyidagi g'oya uchun batafsil rasm prompti yoz. "
                "Uslub: " + style +
                ". G'oya: " + idea
            )

            with st.spinner("Prompt yaratilmoqda..."):
                result = ask_ai(prompt)
                st.markdown("### Tayyor prompt")
                st.write(result)

                st.button(
                    "📋 Promptni nusxalash",
                    on_click=lambda: st.write(
                        "Promptni yuqoridan belgilang va nusxalang."
                    )
                )
        else:
            st.warning("Rasm g'oyasini yozing.")


elif page == "📈 Excel tahlil":
    st.title("📈 JON AI — Excel tahlil")

    uploaded_file = st.file_uploader(
        "Excel yoki CSV fayl yuklang",
        type=["xlsx", "csv"]
    )

    if uploaded_file:
        try:
            if uploaded_file.name.endswith(".csv"):
                df = pd.read_csv(uploaded_file)
            else:
                df = pd.read_excel(uploaded_file)

            st.success("Fayl muvaffaqiyatli ochildi!")
            st.dataframe(df, use_container_width=True)

            st.markdown("### Asosiy ma'lumotlar")
            st.write(df.describe(include="all"))

            question = st.text_input(
                "Jadval haqida savolingiz"
            )

            if st.button("📊 Tahlil qilish"):
                if question:
                    context = df.head(30).to_string()
                    answer = ask_ai(
                        question,
                        "Jadval ma'lumotlari:\n" + context
                    )
                    st.markdown(answer)
                else:
                    st.warning("Savol kiriting.")

        except Exception as e:
            st.error(f"Faylni o'qishda xatolik: {e}")


elif page == "📝 Vazifalar":
    st.title("📝 JON AI — Kunlik vazifalar")

    new_task = st.text_input(
        "Yangi vazifa"
    )

    if st.button("➕ Vazifa qo'shish"):
        if new_task:
            st.session_state.tasks.append(new_task)
            st.rerun()

    st.markdown("### Vazifalar ro'yxati")

    for i, task in enumerate(st.session_state.tasks):
        st.checkbox(
            task,
            key=f"task_{i}"
        )

    if st.button("🗑️ Barcha vazifalarni tozalash"):
        st.session_state.tasks = []
        st.rerun()


elif page == "📷 Kamera":
    st.title("📷 JON AI — Kamera")

    photo = st.camera_input("Suratga olish")

    if photo:
        st.image(
            photo,
            caption="Olingan surat",
            use_container_width=True
        )


elif page == "👤 Personaj":
    st.title("👤 JON AI — Personaj sozlamalari")

    st.text_input(
        "Personaj nomi",
        key="character_name"
    )

    st.selectbox(
        "Tashqi ko'rinish",
        [
            "Hologramma",
            "Robot",
            "Multfilm uslubi",
            "Fantastik"
        ],
        key="character_style"
    )

    avatar = st.file_uploader(
        "Personaj rasmi yuklash",
        type=["png", "jpg", "jpeg"]
    )

    if avatar:
        st.session_state.avatar = avatar.getvalue()

    if st.session_state.avatar:
        st.image(
            st.session_state.avatar,
            caption="Siz tanlagan personaj",
            use_container_width=True
        )
    else:
        show_character()

    st.info(
        "Bu versiyada yuklangan rasm personaj portreti sifatida "
        "ko'rsatiladi. U avtomatik 3D modelga aylantirilmaydi."
    )


elif page == "⚙️ Sozlamalar":
    st.title("⚙️ JON AI — Sozlamalar")

    st.markdown("### Gemini API")

    st.text_input(
        "Gemini API kaliti",
        type="password",
        key="api_key"
    )

    st.caption(
        "API kalitingizni hech kimga yubormang yoki oshkor qilmang."
    )

    if st.button("🔄 API ulanishini tekshirish"):
        if st.session_state.api_key:
            with st.spinner("Tekshirilmoqda..."):
                test = ask_ai("Salom. Faqat OK deb javob ber.")
                if test.startswith("Xatolik:"):
                    st.error(test)
                else:
                    st.success("API javob berdi!")
        else:
            st.warning("API kalitini kiriting.")

    st.divider()
    st.write("Ilova versiyasi:", APP_VERSION)
    st.write("Model:", MODEL)


elif page == "💎 Premium":
    st.title("💎 JON AI Premium")

    st.subheader("Kelajakdagi imkoniyatlar")

    st.write("🚀 Kengaytirilgan AI imkoniyatlari")
    st.write("🎭 Kengaytirilgan personaj sozlamalari")
    st.write("🎙️ Ovozli yordamchi")
    st.write("✨ Maxsus dizayn va effektlar")

    st.info(
        "Premium bo'limi hozircha namoyish uchun. "
        "To'lov tizimi ulanmagan."
    )

st.sidebar.divider()
st.sidebar.caption("JON AI — Your Virtual Assistant")

# ==============================
# 🖼️ JON AI — IMAGE FUNCTION
# ==============================

def get_image(uploaded_file):
    if uploaded_file is None:
        return None

    try:
        image_bytes = uploaded_file.getvalue()

        image_data = {
            "mime_type": uploaded_file.type,
            "data": image_bytes
        }

        return image_data

    except Exception as e:
        st.error(f"Rasmni yuklashda xatolik: {e}")
        return None
        
        
